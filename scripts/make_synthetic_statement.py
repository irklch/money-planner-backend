"""Синтетическая «большая» выписка Альфа-Банка для нагрузочных проверок.

Блок операций обезличенного образца повторяется N раз; шапка и подвал сохраняются.
    python scripts/make_synthetic_statement.py 40 /tmp/alfa_synthetic_x40.xlsx   # ≈ 4 800 расходов
Нужен openpyxl (dev-зависимость). Реальные выписки не использовать.
"""

import sys

import openpyxl

FIXTURE = "tests/fixtures/banks/alfa/alfa_statement_test.xlsx"
HEADER_ROWS = 20  # строки до первой операции, включая заголовок таблицы


def main(copies: int, out: str) -> None:
    src = openpyxl.load_workbook(FIXTURE).active
    rows = [[c.value for c in r] for r in src.iter_rows()]
    end = HEADER_ROWS
    while end < len(rows) and any(v is not None for v in rows[end]):
        end += 1
    head, block, tail = rows[:HEADER_ROWS], rows[HEADER_ROWS:end], rows[end:]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = src.title
    for r in head + block * copies + tail:
        ws.append(r)
    wb.save(out)
    print(f"{len(block)} rows x {copies} -> {out}")


if __name__ == "__main__":
    main(int(sys.argv[1]), sys.argv[2])
