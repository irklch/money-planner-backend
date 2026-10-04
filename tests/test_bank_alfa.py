"""Адаптер Альфа-Банка: образец tests/fixtures/banks/alfa/ + синтетические выписки того же формата.

TEST_SHOP_* в образце вымышленные, поэтому категоризацию проверяем на отдельных синтетических
описаниях с подставным LLM.
"""

import io
import logging
import uuid
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import func, select

from app.ai.provider import set_llm_client
from app.db.models import Expense
from app.db.session import sessionmaker
from app.modules.imports.masking import mask_pii
from app.modules.imports.parsing import registry
from app.modules.imports.parsing.alfa import AlfaBankAdapter
from app.modules.imports.parsing.reader import read_table
from tests.conftest import CAFE, OTHER, PRODUCTS
from tests.test_parse import FakeLLM

FIXTURE = Path(__file__).parent / "fixtures" / "banks" / "alfa" / "alfa_statement_test.xlsx"

# ---------- Независимый расчёт по образцу (openpyxl, без кода адаптера) ----------


@dataclass
class Expected:
    rows: int
    debits: list[Decimal]
    credits: list[Decimal]
    own_debits: list[Decimal]
    own_credits: list[Decimal]
    c2b: list[Decimal]
    p2p_out: list[Decimal]


def independent_count() -> Expected:
    ws = openpyxl.load_workbook(FIXTURE, read_only=True).active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    head = next(i for i, r in enumerate(rows) if r[0] == "Дата операции")
    names = [" ".join(str(c or "").split()) for c in rows[head]]
    col = {n: names.index(n) for n in ("Дата операции", "Описание", "Сумма в валюте счета")}
    e = Expected(0, [], [], [], [], [], [])
    for r in rows[head + 1 :]:
        if not (isinstance(r[0], str) and r[0][:2].isdigit()):
            continue
        e.rows += 1
        amount = Decimal(r[col["Сумма в валюте счета"]].replace(",", "."))
        desc = r[col["Описание"]]
        own = desc.startswith("Внутрибанковский перевод между счетами")
        if amount < 0:
            e.debits.append(-amount)
            (e.own_debits if own else []).append(-amount)
            if "C2B" in desc:
                e.c2b.append(-amount)
            if "Перевод по СБП" in desc:
                e.p2p_out.append(-amount)
        else:
            e.credits.append(amount)
            (e.own_credits if own else []).append(amount)
    return e


EXP = independent_count()


def test_independent_count_of_sample():
    # Значения сверены вручную по образцу; итоговые строки банка в расчёте не используются.
    assert EXP.rows == 161
    assert (len(EXP.debits), sum(EXP.debits)) == (131, Decimal("309294.74"))
    assert (len(EXP.own_debits), sum(EXP.own_debits)) == (11, Decimal("15370.91"))
    assert (len(EXP.credits), sum(EXP.credits)) == (30, Decimal("82290.67"))
    assert (len(EXP.own_credits), sum(EXP.own_credits)) == (19, Decimal("54114.11"))
    assert (len(EXP.c2b), sum(EXP.c2b)) == (4, Decimal("7452.23"))
    assert len(EXP.p2p_out) == 25


# ---------- Синтетические выписки в формате Альфа-Банка ----------

HEADER = ["Дата операции", "Дата проводки", None, "Код", "Категория", *[None] * 6, "Описание"]
HEADER += ["Сумма в\xa0валюте\xa0счета", None, "Статус "]
CLIENT = "ИВАНОВ ИВАН ИВАНОВИЧ"


def op_row(op_date, amount, desc="Тестовая покупка", status="Выполнен", posted=None, code="A1", cat=None):
    cat = cat or "Прочие операции"
    posted = op_date if posted is None else posted
    return [op_date, posted, None, code, cat, *[None] * 6, desc, amount, None, status]


def alfa_xlsx(
    ops,
    *,
    currency="RUR",
    client=CLIENT,
    title="Выписка по счету",
    bank_mark=True,
    header=None,
    lead=6,
) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    for _ in range(lead):
        ws.append([])
    ws.append([title])
    period = "За период с 01.04.2024 по 30.04.2024"
    ws.append(["Номер счета", None, "40817810000000000001", *[None] * 7, period])
    ws.append(["Валюта счета", None, currency, *[None] * 7, "Расходы", None, None, "999,99 RUR"])
    ws.append(["Клиент", None, client])
    ws.append([])
    ws.append(["Операции по счету"])
    ws.append([])
    ws.append(header or HEADER)
    for r in ops:
        ws.append(r)
    ws.append([])
    ws.append([*[None] * 7, "ТЕСТОВЫЙ СОТРУДНИК"])
    if bank_mark:
        bank = "АО «АЛЬФА-БАНК»"
        ws.append([f"(подпись сотрудника {bank})", *[None] * 6, f"(Ф.И.О. сотрудника {bank})"])
    ws.append(["Страница 1 из 1"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


BASIC = [op_row("01.04.2024", "-100,00")]
OWN = (
    f"Внутрибанковский перевод между счетами, {CLIENT}. "
    "Со счёта 40817810000000000001 на счёт 40817810000000000002"
)


def detect(data: bytes, fmt: str = "xlsx"):
    return registry.detect(read_table(data, fmt))


def parse(data: bytes):
    return AlfaBankAdapter().parse(read_table(data, "xlsx"))


async def upload(client, headers, data: bytes, filename="statement.xlsx"):
    return await client.post("/v1/imports/parse", files={"file": (filename, data)}, headers=headers)


# ---------- Распознавание формата ----------


def test_detects_sample():
    assert isinstance(detect(FIXTURE.read_bytes()), AlfaBankAdapter)


def test_detects_shifted_header_and_line_breaks():
    header = list(HEADER)
    header[0], header[12] = "Дата\nоперации", "Сумма  в валюте\nсчёта"
    assert isinstance(detect(alfa_xlsx(BASIC, header=header, lead=15)), AlfaBankAdapter)


@pytest.mark.parametrize(
    "kwargs",
    [
        {"bank_mark": False},  # нет подписи АО «АЛЬФА-БАНК»
        {"title": "Отчёт"},  # нет «Выписка по счету»
        {"currency": "USD"},  # поддержана только рублёвая выписка
        {"header": [c if c != "Код" else None for c in HEADER]},  # нет обязательной колонки
        {"header": ["Дата", "Сумма", "Описание"]},  # любая таблица «Дата/Сумма» — не Альфа
    ],
)
def test_does_not_detect_lookalikes(kwargs):
    assert detect(alfa_xlsx(BASIC, **kwargs)) is None


async def test_csv_is_not_declared_supported(client, headers):
    rows = read_table(FIXTURE.read_bytes(), "xlsx").sheets[0].rows
    csv = "\n".join(";".join(str(c) for c in r) for r in rows).encode("utf-8")
    assert detect(csv, "csv") is None
    r = await upload(client, headers, csv, "alfa.csv")
    assert r.status_code == 422 and r.json()["error"]["code"] == "unknown_bank"


# ---------- Разбор образца ----------


def test_parse_sample_matches_independent_count():
    res = parse(FIXTURE.read_bytes())
    ops, own = res.operations, res.own_transfers
    assert res.unread == []
    assert len(ops) + len(own) == EXP.rows  # служебные строки и итоги шапки не стали операциями
    assert res.period == (date(2024, 4, 1), date(2024, 4, 30))

    all_debits = sorted(-o.amount for o in (*ops, *own) if o.amount < 0)
    assert all_debits == sorted(EXP.debits)
    assert sorted(-o.amount for o in own if o.amount < 0) == sorted(EXP.own_debits)
    assert sorted(o.amount for o in own if o.amount > 0) == sorted(EXP.own_credits)
    expenses = [-o.amount for o in ops if o.amount < 0]
    assert (len(expenses), sum(expenses)) == (120, Decimal("293923.83"))
    assert sum(expenses) == sum(EXP.debits) - sum(EXP.own_debits)
    assert all(isinstance(o.amount, Decimal) for o in ops)


def test_sample_dates_are_operation_dates():
    by_row = {o.row: o for o in parse(FIXTURE.read_bytes()).operations}
    # Строка 29: операция 01.04.2024, проводка 03.04.2024.
    assert by_row[29].date == date(2024, 4, 1)
    # Строка 177: проводка (01.04) раньше операции (28.04). Это аномалия синтетического образца,
    # а не правило реальной выписки; проверяем только, что берётся дата операции.
    assert by_row[177].date == date(2024, 4, 28)
    assert all(type(o.date) is date for o in by_row.values())


def test_sample_statuses():
    res = parse(FIXTURE.read_bytes())
    rows = read_table(FIXTURE.read_bytes(), "xlsx").sheets[0].rows
    statuses = {" ".join(str(rows[o.row - 1][14]).split()) for o in res.operations}
    # «Выполнен» и пустой статус при заполненной дате проводки.
    assert statuses == {"Выполнен", ""}


# ---------- Статусы, суммы и повреждённые строки ----------


def test_amount_formats():
    res = parse(
        alfa_xlsx(
            [
                op_row("01.04.2024", "-1\xa0234,56"),
                op_row("02.04.2024", "-12 345,67"),
                op_row("03.04.2024", -99.9),
                op_row("04.04.2024", "+500,00"),
            ]
        )
    )
    assert [o.amount for o in res.operations] == [
        Decimal("-1234.56"),
        Decimal("-12345.67"),
        Decimal("-99.90"),
        Decimal("500.00"),
    ]


def test_damaged_rows_are_reported_not_imported():
    res = parse(
        alfa_xlsx(
            [
                op_row("01.04.2024", "-100,00"),
                op_row("02.04.2024", "-12,34,56"),  # повреждённая сумма
                op_row("31.02.2024", "-10,00", posted="03.04.2024"),  # несуществующая дата
                op_row("04.04.2024", "-10,00", status="В обработке"),  # неизвестный статус
                op_row("05.04.2024", "-10,00", status="", posted=""),  # без статуса и проводки
                op_row("06.04.2024", None),  # нет суммы
            ]
        )
    )
    assert [o.amount for o in res.operations] == [Decimal("-100.00")]
    rows = [u.row for u in res.unread]
    assert rows == list(range(rows[0], rows[0] + 5))  # номера строк — для диагностики, без содержимого
    assert [u.date for u in res.unread] == [date(2024, 4, d) for d in (2, 3, 4, 5, 6)]


async def test_api_damaged_rows_give_partially_read(client, headers):
    data = alfa_xlsx([op_row("01.04.2024", "-100,00"), op_row("02.04.2024", "abc")])
    r = await upload(client, headers, data)
    assert r.status_code == 200, r.text
    body = r.json()
    assert [o["amount"] for o in body["operations"]] == ["100.00"]
    assert body["warnings"] == [
        {"code": "partially_read", "unreadDates": ["2024-04-02"], "unreadRowCount": 1}
    ]


async def test_api_all_rows_damaged_is_corrupted(client, headers):
    r = await upload(client, headers, alfa_xlsx([op_row("01.04.2024", "abc")]))
    assert r.status_code == 422 and r.json()["error"]["code"] == "corrupted_file"
    assert "abc" not in r.text


def test_own_transfer_requires_statement_owner():
    other = OWN.replace(CLIENT, "ПЕТРОВ ПЁТР ПЕТРОВИЧ")
    res = parse(alfa_xlsx([op_row("01.04.2024", "-100,00", OWN), op_row("02.04.2024", "-50,00", other)]))
    assert [o.amount for o in res.own_transfers] == [Decimal("-100.00")]
    assert [o.amount for o in res.operations] == [Decimal("-50.00")]


# ---------- API: разбор образца и сохранение ----------


async def count_expenses() -> int:
    async with sessionmaker()() as s:
        return await s.scalar(select(func.count()).select_from(Expense))


async def test_api_parse_sample_then_import_idempotent(client, headers):
    r = await upload(client, headers, FIXTURE.read_bytes())
    assert r.status_code == 200, r.text
    body = r.json()
    ops = body["operations"]
    assert body["bank"] == {"code": "alfa", "name": "Альфа-Банк"}
    assert body["period"] == {"from": "2024-04-01", "to": "2024-04-30"}
    assert body["warnings"] == []
    # Поступления без переводов между своими счетами.
    assert body["skippedIncomeCount"] == len(EXP.credits) - len(EXP.own_credits) == 11
    assert len(ops) == 120
    assert sum(Decimal(o["amount"]) for o in ops) == Decimal("293923.83")
    # Переводы между своими счетами исключены, СБП C2B остались покупками.
    assert not any("между счетами" in (o["comment"] or "") for o in ops)
    c2b = [Decimal(o["amount"]) for o in ops if "C2B" in o["comment"]]
    assert sorted(c2b) == sorted(EXP.c2b)
    # Переводы по СБП физлицам по текущему контракту остаются в списке на проверку.
    assert len([o for o in ops if o["comment"].startswith("Категория: Перевод по СБП")]) == 25
    # Без ИИ вымышленные магазины не получают категорий.
    assert {(o["categoryId"], o["categoryStatus"]) for o in ops} == {(None, "unassigned")}

    # Пользователь проверил список и выбрал категорию — сохраняем.
    payload = {
        "expenses": [
            {"date": o["date"], "amount": o["amount"], "categoryId": OTHER, "comment": o["comment"]}
            for o in ops
        ]
    }
    h = {**headers, "Idempotency-Key": str(uuid.uuid4())}
    r1 = await client.post("/v1/expenses/import", json=payload, headers=h)
    assert r1.status_code == 201, r1.text
    assert r1.json() == {"importedCount": 120, "dateFrom": "2024-04-01", "dateTo": "2024-04-30"}
    async with sessionmaker()() as s:
        assert await s.scalar(select(func.sum(Expense.amount))) == Decimal("293923.83")

    r2 = await client.post("/v1/expenses/import", json=payload, headers=h)
    assert r2.status_code == 201 and r2.headers["Idempotent-Replayed"] == "true"
    assert r2.json() == r1.json()
    assert await count_expenses() == 120


# ---------- Категоризация и ПДн ----------

C2B = (
    "Категория: Исходящий платеж QR по СБП C2B. "
    "Платеж A410112 в ПЯТЕРОЧКА 1234 через Систему быстрых платежей."
)
CAFE_OP = "Оплата 220015******1234 КОФЕЙНЯ СИНТЕТИКА MOSCOW RUS"
P2P = (
    "Категория: Перевод по СБП. Перевод B410112 через Систему быстрых платежей на +79161234567 "
    "ИВАН ИВАНОВИЧ И.. Без НДС."
)
INCOME = (
    "Категория: Перевод по СБП. Перевод B410113 через Систему быстрых платежей от +79031112233 "
    "ПЁТР П.. Без НДС."
)


async def test_categorization_with_synthetic_descriptions(client, headers):
    llm = FakeLLM('{"items":[{"i":0,"c":1,"conf":"medium"},{"i":1,"c":0,"conf":"low"}]}')
    set_llm_client(llm)
    data = alfa_xlsx(
        [
            op_row("01.04.2024", "-237,01", C2B, cat="Финансовые операции"),
            op_row("02.04.2024", "-450,50", CAFE_OP, status="", posted="04.04.2024"),
            op_row("03.04.2024", "-1 500,00", P2P),
            op_row("04.04.2024", "-700,00", OWN),
            op_row("05.04.2024", "2 000,00", INCOME),
        ]
    )
    r = await upload(client, headers, data)
    assert r.status_code == 200, r.text
    ops = r.json()["operations"]
    assert [(o["amount"], o["categoryId"], o["categoryStatus"]) for o in ops] == [
        ("237.01", PRODUCTS, "assigned"),  # детерминированное правило по названию сети
        ("450.50", CAFE, "suggested"),
        ("1500.00", None, "unassigned"),
    ]
    assert r.json()["skippedIncomeCount"] == 1

    # В ИИ ушли только маскированные названия расходов: без сумм, дат, телефонов, ФИО, карт и счетов.
    assert len(llm.calls) == 1
    sent = " ".join(m.text for m in llm.calls[0][1])
    for secret in ("9161234567", "ИВАН", "220015", "1234", "450", "2024", "4081781", "9031112233", "ПЁТР"):
        assert secret not in sent, secret
    assert "КОФЕЙНЯ СИНТЕТИКА" in sent


async def test_parse_does_not_log_statement_content(client, headers, caplog):
    caplog.set_level(logging.DEBUG)
    r = await upload(client, headers, FIXTURE.read_bytes())
    assert r.status_code == 200
    logged = " ".join(f"{rec.getMessage()} {rec.__dict__}" for rec in caplog.records)
    for secret in ("TEST_SHOP", "TEST_OPERATION", "ТЕСТОВЫЙ ПОЛЬЗОВАТЕЛЬ", "СЧЁТ_", "237,01", "237.01"):
        assert secret not in logged, secret


@pytest.mark.parametrize(
    ("src", "secrets"),
    [
        (
            "Внутрибанковский перевод между счетами, ИВАНОВ ИВАН ИВАНОВИЧ. "
            "Со счёта 40817810104560012345 на счёт 40817 81010 45600 12346",
            ("ИВАНОВ", "40817810104560012345", "12346"),
        ),
        (
            "Категория: Перевод по СБП. Перевод B4101 через Систему быстрых платежей "
            "на +79161234567. Без НДС.",
            ("9161234567",),
        ),
        (
            "Категория: Перевод по СБП. Перевод B4101 через Систему быстрых платежей от "
            "+7 (903) 111-22-33 Петрова Анна Сергеевна. Без НДС.",
            ("903", "111-22-33", "Петрова", "Анна"),
        ),
        ("Перевод по СБП на 8 916 123 45 67 ИВАН ИВАНОВИЧ И.", ("916", "ИВАН")),
        ("Оплата 220015******1234 PYATEROCHKA 5678 MOSCOW RUS", ("220015", "1234")),
        ("Перевод с карты 427601++++++5678 на карту 2200 1512 3456 7890", ("427601", "5678", "7890")),
        ("Перевод клиенту Альфа-Банка ПЕТРОВ П. по номеру счёта 40817810000000000002", ("ПЕТРОВ", "40817")),
    ],
)
def test_mask_alfa_descriptions(src, secrets):
    masked = mask_pii(src)
    for s in secrets:
        assert s not in masked, (s, masked)


def test_mask_keeps_merchant_in_c2b():
    assert "ПЯТЕРОЧКА" in mask_pii(C2B)


# ---------- Правила первой версии (docs/banks/alfa.md) ----------


def test_rule_1_operation_date_and_posting_date_only_for_completion():
    res = parse(
        alfa_xlsx(
            [
                op_row("01.04.2024", "-10,00", status="", posted="03.04.2024"),
                op_row("02.04.2024", "-20,00", status="", posted=""),
            ]
        )
    )
    assert [(o.date, o.amount) for o in res.operations] == [(date(2024, 4, 1), Decimal("-10.00"))]
    assert [u.date for u in res.unread] == [date(2024, 4, 2)]


async def test_rules_2_to_4_transfers_c2b_and_income(client, headers):
    refund_like = "Возврат покупки TEST_SHOP_A"
    data = alfa_xlsx(
        [
            op_row("01.04.2024", "-1 500,00", P2P),  # перевод физлицу — расход на проверку
            op_row("02.04.2024", "-237,01", C2B, cat="Финансовые операции"),  # покупка СБП C2B
            op_row("03.04.2024", "-700,00", OWN),  # исходящий перевод себе — исключён
            op_row("04.04.2024", "700,00", OWN),  # входящий перевод себе — исключён
            op_row("05.04.2024", "2 000,00", INCOME),  # поступление — пропущено
            op_row("06.04.2024", "300,00", refund_like),  # возврат не угадывается — поступление
        ]
    )
    r = await upload(client, headers, data)
    assert r.status_code == 200, r.text
    body = r.json()
    assert [(o["date"], o["amount"]) for o in body["operations"]] == [
        ("2024-04-01", "1500.00"),
        ("2024-04-02", "237.01"),
    ]
    assert body["skippedIncomeCount"] == 2  # без входящего перевода между своими счетами


async def test_rule_5_other_statuses_are_partially_read(client, headers):
    data = alfa_xlsx(
        [
            op_row("01.04.2024", "-100,00"),
            op_row("02.04.2024", "-50,00", status="В обработке"),
            op_row("03.04.2024", "-60,00", status="Отклонен"),
        ]
    )
    body = (await upload(client, headers, data)).json()
    assert [o["amount"] for o in body["operations"]] == ["100.00"]
    assert body["warnings"] == [
        {"code": "partially_read", "unreadDates": ["2024-04-02", "2024-04-03"], "unreadRowCount": 2}
    ]


@pytest.mark.parametrize("currency", ["USD", "EUR", "CNY"])
async def test_rule_6_only_ruble_accounts(client, headers, currency):
    r = await upload(client, headers, alfa_xlsx(BASIC, currency=currency))
    assert r.status_code == 422 and r.json()["error"]["code"] == "unknown_bank"


async def test_rule_7_bank_operation_types_are_not_categories(client, headers):
    llm = FakeLLM('{"items":[]}')
    set_llm_client(llm)
    data = alfa_xlsx([op_row("01.04.2024", "-100,00", "Оплата ТОВАРЫ СИНТЕТИКА", cat="Супермаркеты")])
    ops = (await upload(client, headers, data)).json()["operations"]
    assert (ops[0]["categoryId"], ops[0]["categoryStatus"]) == (None, "unassigned")
    sent = " ".join(m.text for m in llm.calls[0][1])
    assert "Супермаркеты" not in sent and "ТОВАРЫ СИНТЕТИКА" in sent
